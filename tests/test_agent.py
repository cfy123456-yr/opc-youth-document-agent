import base64
import os
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

from opc_agent.agent import YouthDocumentAgent
from opc_agent.models import AgentRequest, TaskType
from opc_agent.parsers import (
    DOCX_IMAGE_ONLY_WARNING,
    PDF_SCANNED_WARNING,
    ParsedDocument,
    _parse_image,
    _paddle_ocr_engine,
    parse_document,
)


class AgentTests(unittest.TestCase):
    def test_docx_uses_standard_library_when_python_docx_is_missing(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "resume.docx"
            document_xml = """<?xml version="1.0" encoding="UTF-8"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:r><w:t>张三</w:t></w:r></w:p>
    <w:tbl>
      <w:tr>
        <w:tc><w:p><w:r><w:t>技能</w:t></w:r></w:p></w:tc>
        <w:tc><w:p><w:r><w:t>Python</w:t></w:r></w:p></w:tc>
      </w:tr>
    </w:tbl>
  </w:body>
</w:document>
"""
            with zipfile.ZipFile(source, "w") as archive:
                archive.writestr("word/document.xml", document_xml)

            with mock.patch.dict(sys.modules, {"docx": None}):
                parsed = parse_document(source)

            self.assertEqual(parsed.parser, "docx-xml")
            self.assertIn("张三", parsed.text)
            self.assertIn("技能 | Python", parsed.text)
            self.assertEqual(parsed.warnings, [])

    def test_auto_route_and_artifacts(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            agent = YouthDocumentAgent(output_dir=temp_dir)
            run = agent.run(
                AgentRequest(
                    task=TaskType.AUTO,
                    content="请审查这份实习协议中的违约金和竞业限制条款",
                )
            )
            self.assertEqual(run.decision.task, TaskType.CONTRACT)
            self.assertEqual(len(run.output_files), 2)
            for path in run.output_files:
                self.assertTrue(Path(path).exists())

    def test_unknown_task_raises(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            agent = YouthDocumentAgent(output_dir=temp_dir)
            with self.assertRaises(ValueError):
                agent.run(
                    AgentRequest(
                        task=TaskType.AUTO,
                        content="普通文本，没有业务关键词",
                    )
                )

    def test_parser_warnings_are_written_to_report(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "resume.unknown"
            source.write_text(
                "张三\n求职简历\nPython 项目经历",
                encoding="utf-8",
            )
            agent = YouthDocumentAgent(
                output_dir=Path(temp_dir) / "artifacts"
            )
            run = agent.run_file(source, task=TaskType.RESUME)

            self.assertIn(
                "未知文件格式，已按纯文本读取。",
                run.result.warnings,
            )
            report = Path(run.output_files[0]).read_text(encoding="utf-8")
            self.assertIn("## 运行提示", report)
            self.assertIn("未知文件格式", report)

    def test_image_only_docx_reports_ocr_warning(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx 未安装")

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            image = root / "pixel.png"
            image.write_bytes(
                base64.b64decode(
                    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwC"
                    "AAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
                )
            )
            source = root / "image-only.docx"
            document = Document()
            document.add_picture(str(image))
            document.save(source)

            with mock.patch(
                "opc_agent.parsers._parse_image",
                side_effect=RuntimeError(
                    'Image OCR requires: pip install -e ".[ocr]"'
                ),
            ):
                parsed = parse_document(source)

            self.assertEqual(parsed.text.strip(), "")
            self.assertIn(DOCX_IMAGE_ONLY_WARNING, parsed.warnings)

    def test_image_only_docx_uses_ocr_text_when_available(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx 未安装")

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            image = root / "pixel.png"
            image.write_bytes(
                base64.b64decode(
                    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwC"
                    "AAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
                )
            )
            source = root / "image-only.docx"
            document = Document()
            document.add_picture(str(image))
            document.save(source)

            with mock.patch(
                "opc_agent.parsers._parse_image",
                return_value=ParsedDocument(
                    text="张三\n求职简历\nPython 项目经历",
                    source=str(image),
                    parser="paddleocr",
                ),
            ):
                parsed = parse_document(source)

            self.assertIn("张三", parsed.text)
            self.assertEqual(parsed.parser, "python-docx+paddleocr")
            self.assertEqual(parsed.warnings, [])

    def test_scanned_pdf_uses_ocr_text(self):
        try:
            import pypdfium2  # noqa: F401
            from PIL import Image, ImageDraw
        except ImportError:
            self.skipTest("pypdfium2 或 Pillow 未安装")

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            image = Image.new("RGB", (640, 240), "white")
            ImageDraw.Draw(image).text((20, 80), "Resume sample", fill="black")
            source = root / "scanned.pdf"
            image.save(source, "PDF")

            with mock.patch(
                "opc_agent.parsers._parse_image",
                return_value=ParsedDocument(
                    text="张三\n求职简历\nPython 项目经历",
                    source=str(source),
                    parser="paddleocr",
                ),
            ):
                parsed = parse_document(source)

            self.assertIn("张三", parsed.text)
            self.assertEqual(parsed.parser, "pypdf+paddleocr")
            self.assertEqual(parsed.warnings, [])

    def test_scanned_pdf_reports_missing_ocr_dependency(self):
        try:
            import pypdfium2  # noqa: F401
            from PIL import Image
        except ImportError:
            self.skipTest("pypdfium2 或 Pillow 未安装")

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            image = Image.new("RGB", (640, 240), "white")
            source = root / "scanned.pdf"
            image.save(source, "PDF")

            with mock.patch(
                "opc_agent.parsers._parse_image",
                side_effect=RuntimeError(
                    'Image OCR requires: pip install -e ".[ocr]"'
                ),
            ):
                parsed = parse_document(source)

            self.assertIn(PDF_SCANNED_WARNING, parsed.warnings)
            self.assertIn(
                'Image OCR requires: pip install -e ".[ocr]"',
                parsed.warnings,
            )

    def test_paddleocr_modern_result_format(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            image = Path(temp_dir) / "receipt.png"
            image.write_bytes(b"not-a-real-image")
            engine = mock.Mock()
            engine.predict.return_value = [
                {
                    "rec_texts": ["某某酒店", "住宿费 320.00元"],
                    "rec_scores": [0.99, 0.98],
                }
            ]

            with mock.patch(
                "opc_agent.parsers._paddle_ocr_engine",
                return_value=engine,
            ):
                parsed = _parse_image(image)

            self.assertEqual(
                parsed.text,
                "某某酒店\n住宿费 320.00元",
            )
            self.assertEqual(parsed.parser, "paddleocr")
            self.assertEqual(parsed.warnings, [])

    def test_paddleocr_legacy_result_format(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            image = Path(temp_dir) / "resume.png"
            image.write_bytes(b"not-a-real-image")
            engine = mock.Mock(spec=["ocr"])
            engine.ocr.return_value = [
                [
                    [[[0, 0], [1, 0], [1, 1], [0, 1]], ["张三", 0.99]],
                    [[[0, 0], [1, 0], [1, 1], [0, 1]], ["Python 开发", 0.98]],
                ]
            ]

            with mock.patch(
                "opc_agent.parsers._paddle_ocr_engine",
                return_value=engine,
            ):
                parsed = _parse_image(image)

            self.assertEqual(parsed.text, "张三\nPython 开发")
            self.assertEqual(parsed.parser, "paddleocr")

    def test_paddleocr_defaults_to_disabled_onednn(self):
        _paddle_ocr_engine.cache_clear()
        with mock.patch.dict(os.environ, {}, clear=False):
            os.environ.pop("PADDLE_PDX_ENABLE_MKLDNN_BYDEFAULT", None)
            fake_paddleocr = mock.Mock()

            class FakePaddleOCR:
                def __init__(
                    self,
                    use_textline_orientation=True,
                    lang=None,
                ):
                    self.use_textline_orientation = use_textline_orientation
                    self.lang = lang

            fake_paddleocr.PaddleOCR = FakePaddleOCR
            with mock.patch.dict(
                "sys.modules",
                {"paddleocr": fake_paddleocr},
            ):
                _paddle_ocr_engine()

            self.assertEqual(
                os.environ["PADDLE_PDX_ENABLE_MKLDNN_BYDEFAULT"],
                "false",
            )
        _paddle_ocr_engine.cache_clear()

    def test_paddleocr_respects_existing_onednn_setting(self):
        _paddle_ocr_engine.cache_clear()
        with mock.patch.dict(
            os.environ,
            {"PADDLE_PDX_ENABLE_MKLDNN_BYDEFAULT": "true"},
            clear=False,
        ):
            fake_paddleocr = mock.Mock()

            class FakePaddleOCR:
                def __init__(
                    self,
                    use_textline_orientation=True,
                    lang=None,
                ):
                    self.use_textline_orientation = use_textline_orientation
                    self.lang = lang

            fake_paddleocr.PaddleOCR = FakePaddleOCR
            with mock.patch.dict(
                "sys.modules",
                {"paddleocr": fake_paddleocr},
            ):
                _paddle_ocr_engine()

            self.assertEqual(
                os.environ["PADDLE_PDX_ENABLE_MKLDNN_BYDEFAULT"],
                "true",
            )
        _paddle_ocr_engine.cache_clear()


if __name__ == "__main__":
    unittest.main()
