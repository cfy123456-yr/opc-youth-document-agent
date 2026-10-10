from __future__ import annotations

import inspect
import json
import os
import sys
import zipfile
import xml.etree.ElementTree as ET
from collections.abc import Mapping
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from tempfile import TemporaryDirectory


DOCX_IMAGE_ONLY_WARNING = (
    "DOCX 没有可提取的文本层，正文可能是图片，需要 OCR。"
)
PDF_SCANNED_WARNING = "PDF 文本层为空，可能是扫描件，需要 OCR。"
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}
PDF_TEXT_MIN_CHARS = 20
PDF_OCR_MAX_PAGES = 20
WORD_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WORD_TAG = f"{{{WORD_NAMESPACE}}}"
VENDOR_DIR = Path(__file__).resolve().parent.parent / "vendor"


@dataclass(slots=True)
class ParsedDocument:
    text: str
    source: str
    parser: str
    warnings: list[str] = field(default_factory=list)


def parse_document(path: str | Path) -> ParsedDocument:
    file_path = Path(path)
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    suffix = file_path.suffix.lower()
    if suffix in {".txt", ".md", ".json", ".csv"}:
        return ParsedDocument(
            text=read_text(file_path),
            source=str(file_path),
            parser="text",
        )
    if suffix in {".pdf", ".docx", ".pptx", ".xlsx"}:
        markdown_result = _try_markitdown(file_path)
        if markdown_result is not None and markdown_result.text.strip():
            return markdown_result
    if suffix == ".docx":
        return _parse_docx(file_path)
    if suffix == ".pdf":
        return _parse_pdf(file_path)
    if suffix in IMAGE_SUFFIXES:
        return _parse_image(file_path)

    return ParsedDocument(
        text=read_text(file_path),
        source=str(file_path),
        parser="text-fallback",
        warnings=["未知文件格式，已按纯文本读取。"],
    )


def read_text(path: str | Path) -> str:
    file_path = Path(path)
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return file_path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return file_path.read_bytes().decode("utf-8", errors="replace")


def _try_markitdown(path: Path) -> ParsedDocument | None:
    try:
        from markitdown import MarkItDown
    except ImportError:
        return None

    try:
        result = MarkItDown().convert(str(path))
        text = result.text_content
    except Exception:
        return None
    return ParsedDocument(
        text=text,
        source=str(path),
        parser="markitdown",
    )


def _parse_docx(path: Path) -> ParsedDocument:
    try:
        from docx import Document
    except ImportError:
        return _parse_docx_xml(path)

    document = Document(str(path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text for cell in row.cells))
    text = "\n".join(paragraphs)
    parser = "python-docx"
    warnings: list[str] = []
    if not text.strip():
        if _docx_contains_media(path):
            text, ocr_warnings = _ocr_docx_media(path)
            warnings.extend(ocr_warnings)
            if text.strip():
                parser = "python-docx+paddleocr"
        else:
            warnings.append("DOCX 没有可提取的文本。")
    return ParsedDocument(
        text=text,
        source=str(path),
        parser=parser,
        warnings=warnings,
    )


def _parse_docx_xml(path: Path) -> ParsedDocument:
    try:
        with zipfile.ZipFile(path) as archive:
            document_xml = archive.read("word/document.xml")
    except (KeyError, OSError, zipfile.BadZipFile) as exc:
        raise RuntimeError("DOCX 文件结构无效，无法读取正文。") from exc

    try:
        root = ET.fromstring(document_xml)
    except ET.ParseError as exc:
        raise RuntimeError("DOCX 正文 XML 无效，无法读取正文。") from exc

    body = root.find(f"{WORD_TAG}body")
    if body is None:
        raise RuntimeError("DOCX 缺少正文节点，无法读取正文。")

    paragraphs = _extract_docx_body_text(body)
    text = "\n".join(paragraphs)
    parser = "docx-xml"
    warnings: list[str] = []
    if not text.strip():
        if _docx_contains_media(path):
            text, ocr_warnings = _ocr_docx_media(path)
            warnings.extend(ocr_warnings)
            if text.strip():
                parser = "docx-xml+paddleocr"
        else:
            warnings.append("DOCX 没有可提取的文本。")
    return ParsedDocument(
        text=text,
        source=str(path),
        parser=parser,
        warnings=warnings,
    )


def _extract_docx_body_text(body: ET.Element) -> list[str]:
    paragraphs: list[str] = []
    for child in body:
        if child.tag == f"{WORD_TAG}p":
            text = _extract_docx_paragraph_text(child)
            if text:
                paragraphs.append(text)
            continue
        if child.tag != f"{WORD_TAG}tbl":
            continue
        for row in child.findall(f"{WORD_TAG}tr"):
            cells: list[str] = []
            for cell in row.findall(f"{WORD_TAG}tc"):
                cell_text = "\n".join(
                    text
                    for paragraph in cell.findall(f"{WORD_TAG}p")
                    if (text := _extract_docx_paragraph_text(paragraph))
                )
                cells.append(cell_text)
            if any(cells):
                paragraphs.append(" | ".join(cells))
    return paragraphs


def _extract_docx_paragraph_text(paragraph: ET.Element) -> str:
    parts: list[str] = []
    for node in paragraph.iter():
        if node.tag == f"{WORD_TAG}t" and node.text:
            parts.append(node.text)
        elif node.tag == f"{WORD_TAG}tab":
            parts.append("\t")
        elif node.tag in {f"{WORD_TAG}br", f"{WORD_TAG}cr"}:
            parts.append("\n")
    return "".join(parts).strip()


def _docx_contains_media(path: Path) -> bool:
    try:
        with zipfile.ZipFile(path) as archive:
            return any(
                name.startswith("word/media/")
                for name in archive.namelist()
            )
    except (OSError, zipfile.BadZipFile):
        return False


def _ocr_docx_media(path: Path) -> tuple[str, list[str]]:
    with TemporaryDirectory(prefix="opc-docx-ocr-") as temp_dir:
        media_paths = _extract_docx_media(path, Path(temp_dir))
        if not media_paths:
            return "", [DOCX_IMAGE_ONLY_WARNING]

        recognized: list[str] = []
        for media_path in media_paths:
            try:
                result = _parse_image(media_path)
            except RuntimeError as exc:
                return "", [DOCX_IMAGE_ONLY_WARNING, str(exc)]
            if result.text.strip():
                recognized.append(result.text.strip())

    if recognized:
        return "\n".join(recognized), []
    return "", ["OCR 没有识别到文本。"]


def _extract_docx_media(path: Path, target_dir: Path) -> list[Path]:
    extracted: list[Path] = []
    try:
        with zipfile.ZipFile(path) as archive:
            for name in sorted(archive.namelist()):
                if not name.startswith("word/media/"):
                    continue
                suffix = Path(name).suffix.lower()
                if suffix not in IMAGE_SUFFIXES:
                    continue
                destination = target_dir / Path(name).name
                destination.write_bytes(archive.read(name))
                extracted.append(destination)
    except (OSError, zipfile.BadZipFile):
        return []
    return extracted


def _parse_pdf(path: Path) -> ParsedDocument:
    PdfReader = _load_pdf_reader()

    reader = PdfReader(str(path))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    parser = "pypdf"
    warnings: list[str] = []
    if len(text.strip()) < PDF_TEXT_MIN_CHARS:
        try:
            ocr_text, ocr_warnings = _ocr_pdf_pages(path)
        except RuntimeError as exc:
            warnings.extend([PDF_SCANNED_WARNING, str(exc)])
        else:
            if ocr_text.strip():
                text = ocr_text
                parser = "pypdf+paddleocr"
                warnings.extend(ocr_warnings)
            else:
                warnings.extend(
                    [PDF_SCANNED_WARNING, *ocr_warnings, "OCR 没有识别到文本。"]
                )
    return ParsedDocument(
        text=text,
        source=str(path),
        parser=parser,
        warnings=warnings,
    )


def _load_pdf_reader():
    try:
        from pypdf import PdfReader

        return PdfReader
    except ImportError:
        pass

    for wheel_path in sorted(VENDOR_DIR.glob("pypdf-*.whl")):
        path_string = str(wheel_path)
        if path_string not in sys.path:
            sys.path.insert(0, path_string)
        try:
            from pypdf import PdfReader

            return PdfReader
        except ImportError:
            continue
    raise RuntimeError(
        'PDF parsing requires: pip install -e ".[documents]"'
    )


def _ocr_pdf_pages(path: Path) -> tuple[str, list[str]]:
    try:
        import pypdfium2 as pdfium
    except ImportError as exc:
        raise RuntimeError(
            'Scanned PDF OCR requires: pip install -e ".[ocr]"'
        ) from exc

    document = pdfium.PdfDocument(str(path))
    page_count = len(document)
    pages_to_scan = min(page_count, PDF_OCR_MAX_PAGES)
    recognized: list[str] = []
    warnings: list[str] = []
    try:
        with TemporaryDirectory(prefix="opc-pdf-ocr-") as temp_dir:
            target_dir = Path(temp_dir)
            for page_index in range(pages_to_scan):
                page = document[page_index]
                try:
                    image = page.render(scale=2).to_pil()
                    image_path = target_dir / f"page-{page_index + 1}.png"
                    image.save(image_path, format="PNG")
                finally:
                    page.close()

                result = _parse_image(image_path)
                if result.text.strip():
                    recognized.append(result.text.strip())
    finally:
        document.close()

    if page_count > pages_to_scan:
        warnings.append(
            f"PDF 共 {page_count} 页，仅 OCR 前 {pages_to_scan} 页。"
        )
    return "\n\n".join(recognized), warnings


def _parse_image(path: Path) -> ParsedDocument:
    engine = _paddle_ocr_engine()
    predict = getattr(engine, "predict", None)
    if callable(predict):
        result = predict(str(path))
    else:
        result = engine.ocr(str(path))

    lines = _extract_ocr_lines(result)
    return ParsedDocument(
        text="\n".join(lines),
        source=str(path),
        parser="paddleocr",
        warnings=[] if lines else ["OCR 没有识别到文本。"],
    )


@lru_cache(maxsize=1)
def _paddle_ocr_engine():
    # Paddle 3.3 oneDNN can fail on OCR graph attributes unless this is disabled.
    os.environ.setdefault("PADDLE_PDX_ENABLE_MKLDNN_BYDEFAULT", "false")
    try:
        from paddleocr import PaddleOCR
    except ImportError as exc:
        raise RuntimeError(
            'Image OCR requires: pip install -e ".[ocr]"'
        ) from exc

    if "use_textline_orientation" in inspect.signature(PaddleOCR).parameters:
        return PaddleOCR(use_textline_orientation=True, lang="ch")
    else:
        return PaddleOCR(lang="ch")


def _extract_ocr_lines(result: object) -> list[str]:
    if isinstance(result, Mapping):
        pages: list[object] = [result]
    elif isinstance(result, (list, tuple)):
        pages = list(result)
    else:
        pages = [result] if result is not None else []

    lines: list[str] = []
    for page in pages:
        if _is_mapping_like(page):
            lines.extend(_extract_modern_ocr_text(page))
            continue
        for item in page or []:
            try:
                text = str(item[1][0]).strip()
            except (IndexError, TypeError):
                continue
            if text:
                lines.append(text)
    return lines


def _is_mapping_like(value: object) -> bool:
    return isinstance(value, Mapping) or hasattr(value, "get")


def _extract_modern_ocr_text(page: object) -> list[str]:
    data: object = page
    if not isinstance(page, Mapping):
        json_value = getattr(page, "json", None)
        if callable(json_value):
            json_value = json_value()
        if isinstance(json_value, str):
            try:
                json_value = json.loads(json_value)
            except json.JSONDecodeError:
                json_value = None
        if isinstance(json_value, Mapping):
            data = json_value

    if not _is_mapping_like(data):
        return []
    texts = data.get("rec_texts", [])
    if not isinstance(texts, (list, tuple)):
        return []
    return [str(text).strip() for text in texts if str(text).strip()]
